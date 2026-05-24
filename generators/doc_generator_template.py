"""
基于漳州职业模板的Word文档生成器
按照模板格式生成专业的建设方案文档
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from typing import Dict, List
import os


class DocGeneratorTemplate:
    """基于模板的Word文档生成器"""

    def __init__(self, output_path: str = "output/建设方案.docx"):
        self.output_path = output_path
        self.doc = Document()
        self._setup_document()
        self._setup_styles()

    def _setup_document(self):
        """设置文档基本属性"""
        # 设置页面边距（参考模板）
        for section in self.doc.sections:
            section.top_margin = Cm(2.54)
            section.bottom_margin = Cm(2.54)
            section.left_margin = Cm(3.17)
            section.right_margin = Cm(3.17)
            section.header_distance = Cm(1.5)
            section.footer_distance = Cm(1.75)

    def _setup_styles(self):
        """设置文档样式（参考模板）"""
        # 设置默认字体 - 宋体小四
        style = self.doc.styles['Normal']
        font = style.font
        font.name = '宋体'
        font.size = Pt(12)  # 小四
        font.color.rgb = RGBColor(0, 0, 0)

        # 中文字体设置
        element = style.element
        rPr = element.find('.//w:rPr', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
        if rPr is not None:
            rFonts = rPr.find('.//w:rFonts', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
            if rFonts is not None:
                rFonts.set(qn('w:eastAsia'), '宋体')
        
        # 设置段落格式
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.first_line_indent = Pt(24)  # 首行缩进2字符
        
        # 创建标题样式
        self._create_heading_styles()
    
    def _create_heading_styles(self):
        """创建标题样式（参考模板）"""
        # 一级标题（章标题）- 黑体三号
        heading1_style = self.doc.styles.add_style('Heading1Custom', 1)
        heading1_font = heading1_style.font
        heading1_font.name = '黑体'
        heading1_font.size = Pt(16)  # 三号
        heading1_font.bold = True
        heading1_font.color.rgb = RGBColor(0, 0, 0)
        
        heading1_element = heading1_style.element
        rPr = heading1_element.find('.//w:rPr', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
        if rPr is not None:
            rFonts = rPr.find('.//w:rFonts', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
            if rFonts is not None:
                rFonts.set(qn('w:eastAsia'), '黑体')
        
        heading1_style.paragraph_format.space_before = Pt(12)
        heading1_style.paragraph_format.space_after = Pt(6)
        heading1_style.paragraph_format.first_line_indent = Pt(0)
        
        # 二级标题（节标题）- 黑体四号
        heading2_style = self.doc.styles.add_style('Heading2Custom', 1)
        heading2_font = heading2_style.font
        heading2_font.name = '黑体'
        heading2_font.size = Pt(14)  # 四号
        heading2_font.bold = True
        heading2_font.color.rgb = RGBColor(0, 0, 0)
        
        heading2_element = heading2_style.element
        rPr = heading2_element.find('.//w:rPr', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
        if rPr is not None:
            rFonts = rPr.find('.//w:rFonts', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
            if rFonts is not None:
                rFonts.set(qn('w:eastAsia'), '黑体')
        
        heading2_style.paragraph_format.space_before = Pt(10)
        heading2_style.paragraph_format.space_after = Pt(4)
        heading2_style.paragraph_format.first_line_indent = Pt(0)

    def add_cover(self, content):
        """添加封面（参考模板格式）"""
        if isinstance(content, dict):
            school_name = content.get('school_name', '')
            project_names = content.get('project_names', [])
            main_title = content.get('main_title', '建设方案')
            reporting_unit = content.get('reporting_unit', '')
            cooperation_unit = content.get('cooperation_unit', '')
            chinese_enterprise = content.get('chinese_enterprise', '')
            footer_company = content.get('footer_company', '')
            date = content.get('date', '')

            # 顶部留白（少量）
            for _ in range(3):
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 学校名称 - 居中加粗
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if school_name:
                run = p.add_run(school_name)
                run.font.size = Pt(16)
                run.font.bold = True
                run.font.name = '黑体'

            # 空行
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 项目名称 - 居中加粗加大
            for project_name in project_names:
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(project_name)
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.name = '黑体'

            # 空行
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 主标题 - 居中加粗
            if main_title:
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(main_title)
                run.font.size = Pt(16)
                run.font.bold = True
                run.font.name = '黑体'

            # 适中空行使单位信息在页面中下部
            for _ in range(8):
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 底部单位信息 - 居中
            info_items = []
            if reporting_unit:
                info_items.append(f"申报单位：{reporting_unit}")
            if cooperation_unit:
                info_items.append(f"合作单位：{cooperation_unit}")
            if chinese_enterprise:
                info_items.append(f"中资企业：{chinese_enterprise}")
            
            for info in info_items:
                p = self.doc.add_paragraph(info)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.runs[0]
                run.font.size = Pt(12)
                run.font.name = '宋体'

            # 日期和公司
            if date:
                p = self.doc.add_paragraph(date)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.runs[0]
                run.font.size = Pt(12)
                run.font.name = '宋体'

            if footer_company:
                p = self.doc.add_paragraph(footer_company)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.runs[0]
                run.font.size = Pt(12)
                run.font.name = '宋体'

        self.doc.add_page_break()

    def add_section(self, section_data: Dict, is_first_section: bool = False):
        """添加章节（参考模板格式）"""
        try:
            if not section_data or len(section_data) == 0:
                return
            
            # 如果不是第一章，在新章节前添加分页符
            if not is_first_section:
                self.doc.add_page_break()
            
            # 添加章节标题 - 黑体三号
            if 'title' in section_data:
                p = self.doc.add_paragraph(section_data['title'])
                p.style = 'Heading1Custom'

            # 添加内容
            if 'content' in section_data and section_data['content']:
                for subsection in section_data['content']:
                    if not subsection or not isinstance(subsection, dict):
                        continue
                    
                    # 添加子标题 - 黑体四号
                    if 'subtitle' in subsection and subsection['subtitle']:
                        p = self.doc.add_paragraph(subsection['subtitle'])
                        p.style = 'Heading2Custom'

                    # 添加段落 - 宋体小四，首行缩进
                    if 'paragraphs' in subsection and subsection['paragraphs']:
                        for para_text in subsection['paragraphs']:
                            text = para_text if isinstance(para_text, str) else str(para_text)
                            
                            # 检查是否为列表项（以数字+点开头，如"1. "、"2. "等）
                            import re
                            is_list_item = bool(re.match(r'^\d+\.\s', text.strip()))
                            # 检查是否为简短的加粗标题（如学校名、公司名，通常不超过30字且无标点结尾）
                            is_bold_title = len(text.strip()) < 30 and not text.strip().endswith(('。', '；', '：', '.', ';', ':')) and not is_list_item
                            
                            p = self.doc.add_paragraph(text)
                            p.style = 'Normal'
                            
                            # 如果是列表项，使用悬挂缩进使换行文字与第一行对齐
                            if is_list_item:
                                p.paragraph_format.first_line_indent = Pt(0)
                                p.paragraph_format.left_indent = Pt(24)  # 整体缩进2字符
                            # 如果是简短标题，也添加适当缩进
                            elif is_bold_title:
                                p.paragraph_format.first_line_indent = Pt(24)

            # 添加表格
            if 'budget_items' in section_data and section_data['budget_items']:
                self._add_budget_table(section_data['budget_items'])

            if 'content' in section_data and section_data['content']:
                for subsection in section_data['content']:
                    if not subsection or not isinstance(subsection, dict):
                        continue
                    if 'table' in subsection and subsection['table']:
                        table_data = subsection['table']
                        if len(table_data) > 0:
                            first_row = table_data[0]
                            if '完成量' in first_row:
                                self._add_outcomes_table(table_data)
                            elif '费用' in first_row or '一级项目' in first_row:
                                self._add_service_list_table(table_data)
        
        except Exception as e:
            import traceback
            print(f"添加章节时出错: {e}")
            traceback.print_exc()

    def _add_outcomes_table(self, table_data: List[Dict]):
        """添加预期成果表格（参考模板格式）"""
        if not table_data or len(table_data) == 0:
            return
        
        # 表格标题
        p = self.doc.add_paragraph("表5-1 预期成果")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.name = '黑体'

        cols = 4
        table = self.doc.add_table(rows=len(table_data) + 1, cols=cols)
        table.style = 'Table Grid'

        # 表头
        headers = ["序号", "指标", "单位", "完成量"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cell.paragraphs[0].runs[0]
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.name = '黑体'

        # 填充数据
        for i, row_data in enumerate(table_data, 1):
            row = table.rows[i]
            
            row.cells[0].text = str(row_data.get('序号', i))
            row.cells[1].text = str(row_data.get('项目内容', ''))
            row.cells[2].text = str(row_data.get('单位', ''))
            row.cells[3].text = str(row_data.get('完成量', ''))

            for cell in row.cells:
                run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.add_run()
                run.font.size = Pt(10)
                run.font.name = '宋体'
            
            row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_service_list_table(self, table_data: List[Dict]):
        """添加服务清单表格（参考模板格式）"""
        if not table_data or len(table_data) == 0:
            return
        
        p = self.doc.add_paragraph("表7-1 服务清单")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.name = '黑体'

        cols = 6
        table = self.doc.add_table(rows=len(table_data) + 2, cols=cols)
        table.style = 'Table Grid'

        headers = ["序号", "一级项目", "二级任务项", "服务内容", "说明", "费用（万元）"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cell.paragraphs[0].runs[0]
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.name = '黑体'

        total_fee = 0
        for i, row_data in enumerate(table_data, 1):
            row = table.rows[i]
            
            row.cells[0].text = str(row_data.get('序号', i))
            row.cells[1].text = str(row_data.get('一级项目', ''))
            row.cells[2].text = str(row_data.get('二级任务项', ''))
            
            service_content = row_data.get('服务内容', '')
            if isinstance(service_content, list):
                service_content = '\n'.join(service_content)
            row.cells[3].text = str(service_content)
            
            row.cells[4].text = str(row_data.get('说明', ''))
            fee_str = str(row_data.get('费用', '0'))
            row.cells[5].text = fee_str

            for cell in row.cells:
                run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.add_run()
                run.font.size = Pt(9)
                run.font.name = '宋体'
            
            try:
                fee_num = float(''.join(filter(lambda c: c.isdigit() or c == '.', fee_str)))
                total_fee += fee_num
            except (ValueError, TypeError):
                pass

        # 合计行
        total_row = table.rows[len(table_data) + 1]
        total_row.cells[0].text = "合计"
        total_row.cells[0].merge(total_row.cells[4])
        total_row.cells[5].text = f"{total_fee:.2f}"
        
        for cell in total_row.cells:
            run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.add_run()
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.name = '黑体'

    def _add_budget_table(self, budget_items: List[Dict]):
        """添加预算表格"""
        p = self.doc.add_paragraph("表6-1 项目预算明细表")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.name = '黑体'

        table = self.doc.add_table(rows=len(budget_items) + 1, cols=5)
        table.style = 'Table Grid'

        headers = ["序号", "项目名称", "数量", "单价(元)", "金额(元)"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cell.paragraphs[0].runs[0]
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.name = '黑体'

        total_amount = 0
        for i, item in enumerate(budget_items, 1):
            row = table.rows[i]
            row.cells[0].text = str(i)
            row.cells[1].text = item.get('item', '')
            row.cells[2].text = str(item.get('quantity', ''))
            row.cells[3].text = f"{item.get('price', 0):,.0f}"
            amount = item.get('quantity', 0) * item.get('price', 0)
            row.cells[4].text = f"{amount:,.0f}"
            total_amount += amount

            for cell in row.cells:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cell.paragraphs[0].runs[0]
                run.font.size = Pt(10)
                run.font.name = '宋体'

        total_row = table.add_row()
        total_row.cells[0].text = "合计"
        total_row.cells[0].merge(total_row.cells[2])
        total_row.cells[3].text = f"{total_amount:,.0f}"
        total_row.cells[3].merge(total_row.cells[4])

        for cell in total_row.cells:
            run = cell.paragraphs[0].runs[0]
            run.font.bold = True
            run.font.name = '黑体'

        p = self.doc.add_paragraph(f"项目总投资：人民币 {total_amount:,.0f} 元（大写：{self._number_to_chinese(total_amount)}）")
        p.paragraph_format.first_line_indent = Pt(24)
        run = p.runs[0]
        run.font.size = Pt(12)

    def _number_to_chinese(self, num: int) -> str:
        """将数字转换为中文大写"""
        chinese_nums = ['零', '壹', '贰', '叁', '肆', '伍', '陆', '柒', '捌', '玖']
        chinese_units = ['', '拾', '佰', '仟']
        chinese_big_units = ['', '万', '亿']

        num = int(num)
        if num == 0:
            return "零元整"

        result = ""
        unit_idx = 0

        while num > 0:
            section = num % 10000
            num //= 10000

            if section > 0:
                section_str = ""
                section_unit_idx = 0

                while section > 0:
                    digit = int(section % 10)
                    section //= 10

                    if digit > 0:
                        section_str = chinese_nums[digit] + chinese_units[section_unit_idx] + section_str
                    elif section_str and section_str[0] != '零':
                        section_str = '零' + section_str

                    section_unit_idx += 1

                result = section_str + chinese_big_units[unit_idx] + result

            unit_idx += 1

        return result + "元整"

    def add_table_of_contents(self):
        """添加目录占位符"""
        p = self.doc.add_paragraph("目 录")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.name = '黑体'

        self.doc.add_paragraph("\n（此处为目录占位符，请在Word中使用'引用->目录'功能生成正式目录）\n")
        
        # 目录后添加分页符，使第一章从新页面开始
        self.doc.add_page_break()

    def generate(self, sections: Dict) -> str:
        """生成完整文档"""
        if 'cover' in sections:
            self.add_cover(sections['cover'])

        self.add_table_of_contents()

        section_keys = ['section1', 'section2', 'section3', 'section4',
                       'section5', 'section6', 'section7']

        is_first_section = True
        for key in section_keys:
            if key in sections and sections[key]:
                self.add_section(sections[key], is_first_section=is_first_section)
                is_first_section = False

        os.makedirs(os.path.dirname(self.output_path), exist_ok=True)
        self.doc.save(self.output_path)

        return self.output_path
