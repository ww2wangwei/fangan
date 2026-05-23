"""
Excel报价单解析器
解析厚溥国际职业教育国际化产品报价单
"""

import openpyxl
import json
import os
from typing import Dict, List, Optional


class ExcelParser:
    """Excel报价单解析器"""

    def __init__(self, file_path: str):
        self.file_path = file_path
        self.workbook = None
        self.data = {}

    def load(self):
        """加载Excel文件"""
        if not os.path.exists(self.file_path):
            raise FileNotFoundError(f"文件不存在: {self.file_path}")

        self.workbook = openpyxl.load_workbook(self.file_path)
        return self

    def parse_all(self) -> Dict:
        """解析所有工作表"""
        if not self.workbook:
            self.load()

        self.data = {
            "product_list": self._parse_product_list(),
            "india_visit": self._parse_india_visit(),
            "video_resources": self._parse_video_resources(),
            "product_catalog": self._parse_product_catalog(),
            "product_summary": self._parse_product_summary(),
        }

        return self.data

    def _parse_sheet_by_name(self, sheet_name: str, max_rows: Optional[int] = None) -> List[Dict]:
        """根据工作表名称解析数据"""
        if sheet_name not in self.workbook.sheetnames:
            return []

        ws = self.workbook[sheet_name]
        headers = [cell.value for cell in ws[1]]
        data = []

        max_row = min(max_rows, ws.max_row) if max_rows else ws.max_row

        for row in ws.iter_rows(min_row=2, max_row=max_row, values_only=True):
            row_dict = {}
            for i, (header, value) in enumerate(zip(headers, row)):
                if header and value is not None:
                    row_dict[str(header).strip()] = str(value).strip() if value else ""
            if any(row_dict.values()):
                data.append(row_dict)

        return data

    def _parse_product_list(self) -> List[Dict]:
        """解析国际化产品清单主表"""
        sheet_name = self.workbook.sheetnames[0]  # 第一个工作表
        ws = self.workbook[sheet_name]

        headers = [cell.value for cell in ws[1]]
        products = []

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
            product = {}
            for i, (header, value) in enumerate(zip(headers, row)):
                if header and value is not None:
                    key = str(header).strip()
                    product[key] = str(value).strip() if value else ""

            # 只保留有有效数据的行(序号和产品名称)
            seq_no = product.get("序号", "")
            product_name = product.get("单项产品", "")
            
            if seq_no and product_name:
                # 提取关键信息（使用正确的列名）
                product['display_name'] = product_name
                product['price'] = product.get("单价", "0")  # 单价
                product['unit'] = product.get("单位", "项")
                product['duration'] = product.get("周期", "")
                # 保留详细内容字段
                product['service_content'] = product.get("服务内容", "")
                product['technical_params'] = product.get("技术参数（改）", "")
                product['evidence_materials'] = product.get("佐证材料/结果清单", "")
                product['category'] = self._categorize_product(product_name)
                products.append(product)

        return products

    def _categorize_product(self, product_name: str) -> str:
        """根据产品名称分类"""
        categories = {
            "院校对接": ["海外开发学校", "海外院校对接", "出访"],
            "专业标准": ["专业标准"],
            "课程标准": ["课程标准"],
            "仪式活动": ["开工仪式", "揭牌仪式"],
            "空间建设": ["教学场所", "实训基地"],
            "师资培训": ["来华留学", "送教上门", "师资培训"],
            "学生培养": ["学历教育", "短期培训"],
            "资源建设": ["视频资源", "教学资源"],
        }

        for category, keywords in categories.items():
            for keyword in keywords:
                if keyword in product_name:
                    return category
        return "其他"

    def _parse_india_visit(self) -> List[Dict]:
        """解析印度出访项目费用"""
        sheet_name = "印度出访项目费用清单（2026版）"
        if sheet_name not in self.workbook.sheetnames:
            return []
        return self._parse_sheet_by_name(sheet_name)

    def _parse_video_resources(self) -> List[Dict]:
        """解析视频资源成本清单"""
        sheet_name = "视频资源成本清单(2026版)"
        if sheet_name not in self.workbook.sheetnames:
            return []
        return self._parse_sheet_by_name(sheet_name)

    def _parse_product_catalog(self) -> List[Dict]:
        """解析国际产品清单"""
        sheet_name = "国际产品清单"
        if sheet_name not in self.workbook.sheetnames:
            return []
        return self._parse_sheet_by_name(sheet_name)

    def _parse_product_summary(self) -> List[Dict]:
        """解析产品汇总表"""
        sheet_name = "产品汇总表"
        if sheet_name not in self.workbook.sheetnames:
            return []
        return self._parse_sheet_by_name(sheet_name)

    def get_products_by_category(self, category: str) -> List[Dict]:
        """按类别获取产品"""
        if not self.data:
            self.parse_all()

        products = self.data.get("product_list", [])
        filtered = []

        for product in products:
            if product.get('category') == category:
                filtered.append(product)

        return filtered

    def get_all_products_for_selection(self) -> List[Dict]:
        """获取所有可选的产品列表(用于建设内容选择)"""
        if not self.data:
            self.parse_all()

        products = self.data.get("product_list", [])
        result = []

        for product in products:
            result.append({
                'id': product.get('序号', ''),
                'name': product.get('单项产品', ''),  # 使用正确的列名
                'service_content': product.get('服务内容', ''),  # 服务内容
                'technical_params': product.get('技术参数（改）', ''),  # 技术参数
                'evidence_materials': product.get('佐证材料/结果清单', ''),  # 佐证材料/结果清单
                'price': product.get('单价', '0'),  # 单价
                'unit': product.get('单位', '项'),
                'category': product.get('category', '其他'),
                'duration': product.get('周期', ''),
                # 兼容旧字段
                'description': product.get('服务内容', '')[:100] if product.get('服务内容') else '',
            })

        return result

    def get_price_range(self, min_price: float = 0, max_price: float = float('inf')) -> List[Dict]:
        """获取价格范围内的产品"""
        if not self.data:
            self.parse_all()

        products = self.data.get("product_list", [])
        filtered = []

        for product in products:
            try:
                price = float(product.get("单价", 0) or 0)
                if min_price <= price <= max_price:
                    filtered.append(product)
            except (ValueError, TypeError):
                continue

        return filtered

    def export_to_json(self, output_path: str = "data/cache.json"):
        """导出数据为JSON缓存"""
        if not self.data:
            self.parse_all()

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(self.data, f, ensure_ascii=False, indent=2)

        return output_path


def main():
    """测试解析器"""
    parser = ExcelParser("厚溥国际职业教育国际化产品报价单2026版(2.0).xlsx")
    data = parser.parse_all()

    print("=" * 60)
    print("Excel报价单解析结果")
    print("=" * 60)

    for sheet_name, items in data.items():
        print(f"\n{sheet_name}: {len(items)} 条记录")
        if items:
            print(f"示例数据: {json.dumps(items[0], ensure_ascii=False, indent=2)[:200]}...")

    # 导出缓存
    cache_file = parser.export_to_json()
    print(f"\n缓存已导出至: {cache_file}")


if __name__ == "__main__":
    main()
